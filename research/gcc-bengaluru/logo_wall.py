"""Standalone workbook for the companies on a sponsor logo wall the user shared
(1 Oct 2026), their HubSpot status, and senior contacts found in Lusha for the ones
Elets has not tapped yet."""
import json
from pathlib import Path

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from outreach import OWNERS, STATUSES

RESULT = Path(__file__).with_name("logo_wall_result.json")
COMPANY_URL = "https://app.hubspot.com/contacts/147308736/record/0-2/{}"


def build_logo_wall(wb, write_table):
    if not RESULT.exists():
        return []
    data = json.loads(RESULT.read_text())
    companies = {c["company"]: c for c in data.get("companies", [])}
    people = {}
    for p in data.get("contacts", []):
        if p.get("name"):
            people.setdefault(p["company"], []).append(p)

    order = {"Untapped": 0, "In CRM, cold": 1, "Tapped": 2}
    rows = []
    for name in sorted(companies, key=lambda n: (order.get(companies[n].get("hubspot_status"), 3), n.lower())):
        c = companies[name]
        status = c.get("hubspot_status", "")
        owner = OWNERS.get(str(c.get("hubspot_owner_id") or ""), str(c.get("hubspot_owner_id") or ""))
        link = COMPANY_URL.format(c["hubspot_company_id"]) if c.get("hubspot_company_id") else ""
        base = [status, name, c.get("what_they_do", ""), c.get("domain", "")]
        crm = [owner, c.get("hubspot_contacts", ""), (c.get("hubspot_last_contacted") or "")[:10], link]
        found = people.get(name, [])
        if not found:
            note = "Already tapped - no Lusha search" if status == "Tapped" else "No senior India contact in Lusha"
            rows.append(base + ["", "", "", "", "", "", ""] + crm + ["Not contacted", "", c.get("notes") or note])
        for p in found:
            rows.append(base + [p.get("name", ""), p.get("title", ""), p.get("role_group", ""), p.get("email", ""),
                                p.get("phone", ""), p.get("linkedin", ""), p.get("location", "")] + crm +
                        ["Not contacted", "", p.get("note", "")])

    ws = wb.create_sheet("Logo Wall Prospects", 0)
    headers = ["HubSpot status", "Company", "What they do", "Website", "Contact name", "Designation",
               "Role group", "Email", "Phone", "LinkedIn", "Location", "HubSpot owner", "HubSpot contacts",
               "HubSpot last contacted", "HubSpot link", "Status", "Elets owner", "Notes"]
    widths = [14, 26, 40, 22, 24, 34, 12, 32, 18, 30, 18, 20, 10, 14, 30, 16, 14, 40]
    write_table(ws, "LogoWall", headers, rows, widths, link_cols=(10, 15))
    fills = {"Untapped": "FCE4D6", "In CRM, cold": "FFF2CC", "Tapped": "E2EFDA"}
    for i in range(2, ws.max_row + 1):
        cell = ws.cell(i, 1)
        if cell.value in fills:
            cell.fill = PatternFill("solid", start_color=fills[cell.value])
            cell.font = Font(name="Arial", size=10, bold=True)
        for col in range(16, 19):
            ws.cell(i, col).fill = PatternFill("solid", start_color="FFFFF2")
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"P2:P{ws.max_row}")
    ws.freeze_panes = "C2"
    return rows


OUT = Path(__file__).with_name("Logo_Wall_Prospects.xlsx")


def main():
    from openpyxl import Workbook
    from build_gcc_sponsors import write_table
    wb = Workbook()
    default = wb.active
    rows = build_logo_wall(wb, write_table)
    wb.remove(default)
    notes = wb.create_sheet("Notes")
    for i, t in enumerate([
            "Logo wall prospects - 35 companies from the sponsor logo wall shared on 1 Oct 2026.",
            "HubSpot status: Untapped = not in HubSpot; In CRM, cold = in HubSpot but never contacted / not in 90 days; "
            "Tapped = contacted in the last 90 days (no Lusha search run for these).",
            "Contacts are senior Marketing / Sales / Leadership people in India from Lusha (about 100 credits used). "
            "Rows marked CHECK in Notes may have moved company - verify before outreach.",
            "Not identified from the image: a black teardrop-icon logo with no name, and a 'Raw Data to AI' wordmark."], 1):
        notes.cell(i, 1, t)
    notes.column_dimensions["A"].width = 130
    wb.save(OUT)
    print(f"{OUT}: {len(rows)} rows")


if __name__ == "__main__":
    main()
