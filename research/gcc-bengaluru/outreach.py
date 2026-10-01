"""Turn the sponsor data into a ranked outreach list.

Categories, contact roles and pitch angles are editorial judgements, not
researched facts; the Outreach sheet says so. Priority score:
  best tier (Presenting/Platinum-level 5, Gold/Strategic 3, Silver/Associate 2, other 1)
  + 3 for every extra event sponsored
  + 2 if they sponsored a 2026 event
  - 2 if every row for the company is an inferred (unconfirmed) one
"""
import json
from pathlib import Path

from openpyxl.worksheet.datavalidation import DataValidation

ENRICH_FILE = Path(__file__).with_name("enrichment.json")


def load_enrichment():
    """Free web enrichment (company site, public contact) keyed by company name."""
    if not ENRICH_FILE.exists():
        return {}
    return {r["company"]: r for r in json.loads(ENRICH_FILE.read_text())}

CATEGORY = {}
for cat, names in {
    "IT services & engineering": [
        "Accenture", "Accion Labs", "Akkodis", "Ascendion", "Birlasoft", "C5i", "Cyient", "EPAM Systems",
        "GlobalLogic", "HCLTech", "Hexaware Technologies", "Infosys", "Tech Mahindra", "Quest Global",
        "Publicis Sapient", "R Systems", "Parkar", "NeoSOFT", "Tavant Technologies", "STL Digital",
        "Incture Technologies", "Paramatrix Technologies", "Avekshaa Technologies",
        "Sigma AVIT Technology Solutions", "Power Bridge Systems", "Sheeltron Digital Systems", "E-Solutions",
        "VentureSoft", "World Wide Technology", "TransformEdge Enterprise Services", "Technogen India",
        "Genpact", "EdgeVerve", "Ideya Labs"],
    "Software, SaaS & AI": [
        "BMC Software", "Commvault", "CoreStack", "DevRev", "Draup", "Freshworks", "monday.com", "Planview",
        "vSaaS Global", "OutSystems", "Progress", "ProHance", "Bluecopa", "CashFlo", "Scry AI", "Snowflake",
        "Google Cloud", "Intellect Design Arena", "OrbitShift.AI", "Simple.AI", "Hosted.AI", "Axentia.AI",
        "Avantra", "Zeko AI", "Ceipal", "Kovaion Consulting", "pCloudy", "Thomson Reuters", "Equiniti India"],
    "Cybersecurity & risk": ["UpGuard", "ColorTokens", "63SATS Cybertech", "AuthBridge"],
    "Hardware & infrastructure": [
        "AMD", "Lenovo", "Dell Technologies", "Hewlett Packard Enterprise", "IBM", "Jabra", "MAXHUB",
        "Brilyant"],
    "Consulting, audit & advisory": [
        "Deloitte", "EY", "PwC India", "Kirtane & Pandit Consulting", "Baker Tilly One India", "CBIZ India",
        "Weaver and Tidwell India", "Dhruva Advisors", "BINDZ Consulting", "Orane Consulting",
        "HSAG Consulting", "Fieldmark Advisory", "Mancer Consulting Services", "Amicorp", "ANSR", "Lex Visas"],
    "Real estate & workspace": [
        "Intellion Offices by Tata Realty", "JLL", "Regalia Business Parks (Hiranandani)", "Awfis",
        "WeWork India", "Table Space", "Flipspaces", "Haworth"],
    "Talent, staffing & learning": [
        "Collabera", "Careernet", "Quess Corp", "Spectrum Talent Management", "Orcapod", "Xpheno",
        "TEKsystems", "Cognixia", "Udemy Business", "upGrad", "LinkedIn"],
    "Employee benefits, wellness & services": [
        "ADP", "Benepik", "cult.fit", "Zyla Health", "Silver Oak Health", "Klay Preschools & Daycare",
        "Silver Brook Learning Center", "Tortoise", "Shework.in", "Elior India", "Cateworld",
        "Uber for Business", "Giriraj Mobility"],
    "BFSI & fintech": ["State Bank of India", "RBL Bank", "IDFC FIRST Bank", "Skydo"],
    "Enterprise running a GCC": ["NatWest Group", "Sony", "Woodside Energy"],
    "Energy & sustainability": ["Fourth Partner Energy"],
    "Govt / ecosystem body (partner, not sponsor)": [
        "Govt of Karnataka (DoITBT)", "Karnataka Digital Economy Mission (KDEM)", "Invest Punjab",
        "Invest UP", "GIFT City", "K-Tech Innovation Hub (IIIT Dharwad Research Park)",
        "Indian Chamber of International Business (ICIB)"],
    "Media & community (partner, not sponsor)": [
        "Prop News Time", "ISAIL.in", "Storify News", "PharmiVon Media LLP", "SSON Research & Analytics",
        "The Sunny Shah Show", "Times Techies (Times of India)", "Koot", "Kaizzen"],
    "Academic (partner, not sponsor)": ["Alliance University", "BITSoM"],
}.items():
    for n in names:
        CATEGORY[n] = cat
# Resolved by the free web enrichment (see enrichment.json "what_they_do").
CATEGORY.update({
    "Everpure": "Hardware & infrastructure", "Everforth": "IT services & engineering",
    "Zertain India": "IT services & engineering", "3R Infotech": "IT services & engineering",
    "Infogini Consulting": "IT services & engineering", "RST Solutions": "IT services & engineering",
    "Syncortex": "IT services & engineering", "Vidushi Infotech": "IT services & engineering",
    "FlexTecs": "Software, SaaS & AI", "Mandaala": "Employee benefits, wellness & services",
    "Famli": "Employee benefits, wellness & services", "Peregrine": "Employee benefits, wellness & services",
    "Ingenious": "Real estate & workspace", "EduTech": "Talent, staffing & learning",
    "ACT Enterprise": "Hardware & infrastructure",
})
OTHER = "Other - check what they do"
NON_SPONSOR = ("partner, not sponsor",)

# (who to contact, pitch angle) per category
ANGLE = {
    "IT services & engineering": (
        "CMO; Head of Marketing / Events; GCC or Enterprise Sales Head (India)",
        "GCC heads are the buyers of build-operate-transfer and engineering services - thought-leadership slot + CXO roundtable."),
    "Software, SaaS & AI": (
        "Head of Marketing India/APJ; Field/Event Marketing Manager; Country Head",
        "Demand generation with GCC tech and ops leaders - product demo booth, AI session, qualified lead list."),
    "Cybersecurity & risk": (
        "Marketing Head India; Regional Sales Director; CISO-community lead",
        "GCC CISOs and risk heads in one room - security panel or closed-door CISO roundtable."),
    "Hardware & infrastructure": (
        "Enterprise/Commercial Marketing Head India; Partner & Alliances Marketing",
        "GCC build-outs and AI infrastructure refresh - platinum branding plus experience zone."),
    "Consulting, audit & advisory": (
        "GCC Practice Leader / Partner; Markets & BD Head; Brand & Marketing Head",
        "Knowledge-partner slot: co-branded GCC report or keynote on GCC setup, tax and operating models."),
    "Real estate & workspace": (
        "Head of Leasing / Business Development; Marketing Head",
        "Occupiers and GCC site selectors - workspace partner branding and site-visit tie-in."),
    "Talent, staffing & learning": (
        "Head of Marketing; Enterprise / GCC Sales Head; Head of Partnerships",
        "GCC hiring and skilling leaders (CHROs, talent heads) - talent-track session and branded lounge."),
    "Employee benefits, wellness & services": (
        "Head of Corporate Sales / Partnerships; Marketing Head",
        "GCC HR and admin decision-makers - exhibitor space, delegate kit, CHRO networking dinner."),
    "BFSI & fintech": (
        "Head of Corporate / Transaction Banking Marketing; Brand Head",
        "Banking partner for GCC set-ups - corporate account and payments sponsorship."),
    "Enterprise running a GCC": (
        "GCC Head / Site Leader; Employer Branding or Talent Brand Lead",
        "Employer-branding and hiring visibility - speaker slot plus employer-brand partnership."),
    "Energy & sustainability": (
        "Head of C&I Sales; Marketing Head",
        "GCC sustainability and RE100 goals - sustainability-partner slot."),
    OTHER: ("Marketing Head; Founder / Business Head", "Check the company first, then pick the angle."),
}
for c in CATEGORY.values():
    ANGLE.setdefault(c, ("Partnerships / Alliances lead", "In-kind partner (ecosystem, media or academic); not a paid-sponsor lead."))

TOP = ("presenting", "co-powered", "platinum", "premium", "global partner", "host state", "lead sponsor",
       "co-presenting", "thought leadership")
MID = ("gold", "strategic", "knowledge", "state partner", "co-host")
LOW = ("silver", "associate", "business orchestration", "co-partner", "digital transformation", "agentic",
       "cxo", "real estate", "talent", "workspace", "banking", "academic", "ecosystem", "supporting",
       "exhibit partner", "ai finance")


def tier_points(tier):
    t = tier.lower()
    for pts, keys in ((5, TOP), (3, MID), (2, LOW)):
        if any(k in t for k in keys):
            return pts
    return 1


STATUSES = ["Not contacted", "Contacted", "In conversation", "Proposal sent", "Won", "Lost", "Not a fit"]


CRM_FILE = Path(__file__).with_name("crm_companies.json")
CRM_URL = "https://app.hubspot.com/contacts/147308736/record/0-2/{}"
OWNERS = {"31264458": "Shivam Pathania", "31264460": "Swati Bhattacharya", "31264463": "Vipul Jain",
          "31264465": "Ishan (inactive)", "31264466": "Akansha Pal", "31381972": "Krishna Kumar Singh",
          "32775033": "Mudit Sharma", "32831970": "Sejal Joshi", "33026142": "Abhay Malhotra",
          "31115589": "Lakshya Kapoor (inactive)", "31381977": "Lakshya Singh (inactive)",
          "32092249": "Anubhav Kumar Dwivedi", "32092253": "Ronak Tiwari", "32092254": "Arpana Singh",
          "32110903": "Anuj Sharma (inactive)", "32881438": "Shreya Kumari", "33026141": "Ayush Agarwal",
          "33183397": "Md. Iqbal", "33609425": "Mohammad Suhail Khan", "33609433": "Ragni Nayak",
          "34032745": "Rinki Verma", "34807069": "Litishia Raina", "34870449": "Sakshi Raj",
          "34870452": "Priyanshi Agrawal", "35025264": "Prem Ranjan Das", "35343623": "Priya Singh",
          "36196623": "Anurag Lakha", "37918715": "Teresa Smrutirekha", "85465367": "Anurag Gupta"}
TODAY = "2026-10-01"


def load_crm():
    """HubSpot company matches (snapshot 1 Oct 2026), merged per sponsor company."""
    if not CRM_FILE.exists():
        return {}
    out = {}
    for name, cid, domain, owner, stage, n_contacts, n_deals, last in json.loads(CRM_FILE.read_text()):
        c = out.setdefault(name, {"ids": [], "owners": [], "stage": "lead", "contacts": 0, "deals": 0, "last": ""})
        c["ids"].append(cid)
        if owner and OWNERS.get(owner) not in c["owners"]:
            c["owners"].append(OWNERS.get(owner, owner))
        if stage == "customer":
            c["stage"] = "customer"
        c["contacts"] += n_contacts
        c["deals"] += n_deals
        c["last"] = max(c["last"], last)
    return out


def days_since(d):
    from datetime import date
    return (date.fromisoformat(TODAY) - date.fromisoformat(d)).days if d else None


def next_action(crm, has_contact):
    if not crm:
        return "New - add to CRM and reach out" + ("" if has_contact else " (find contact first)")
    if crm["stage"] == "customer":
        return "Existing customer - pitch via account owner"
    d = days_since(crm["last"])
    if d is not None and d <= 14:
        return "Active in CRM - coordinate with owner before contacting"
    if d is None:
        return "In CRM, never contacted - start outreach"
    return f"In CRM, cold ({d} days) - re-engage"


# Companies the user asked to leave out of the outreach list (1 Oct 2026).
EXCLUDED_COMPANIES = {"Infosys", "PwC India", "Accenture", "State Bank of India", "Freshworks",
                      "IDFC FIRST Bank", "RBL Bank"}


def build_outreach(wb, events, sponsors, canon, write_table):
    ev = {e[0]: e for e in events}
    by_co = {}
    for eid, tier, name, ver in sponsors:
        by_co.setdefault(canon.get(name, name), []).append((eid, tier, ver))

    enrich = load_enrichment()
    crm_all = load_crm()
    rows = []
    for co, hits in by_co.items():
        if co in EXCLUDED_COMPANIES:
            continue
        e = enrich.get(co, {})
        g = lambda k: (e.get(k) or "").strip()
        src = " | ".join(x for x in (g("contact_source_url"), g("public_email_source"), g("public_phone_source")) if x)
        cat = CATEGORY.get(co, OTHER)
        eids = sorted({h[0] for h in hits})
        best = max(hits, key=lambda h: tier_points(h[1]))
        in_2026 = any("2026" in ev[e][4] for e in eids)
        all_inferred = all(h[2].startswith("Inferred") for h in hits)
        score = tier_points(best[1]) + 3 * (len(eids) - 1) + (2 if in_2026 else 0) - (2 if all_inferred else 0)
        if any(k in cat for k in NON_SPONSOR):
            prio = "Partner only"
        else:
            prio = "A - Hot" if score >= 7 else "B - Warm" if score >= 4 else "C - Nurture"
        history = "; ".join(f"{ev[e][1]} ({ev[e][4]}) - {t}" for e, t, _ in sorted(hits))
        latest = max(eids, key=lambda e: ("2026" in ev[e][4], e))
        who, angle = ANGLE[cat]
        crm = crm_all.get(co)
        action = next_action(crm, bool(g("contact_name"))) if prio != "Partner only" else "In-kind partner - not a sales lead"
        rows.append([
            prio, score, co, cat, g("what_they_do"), g("website"), g("india_hq_city"),
            len(eids), best[1], ev[latest][4], history, "Yes" if all_inferred else "",
            "Yes" if crm else "No",
            ", ".join(crm["owners"]) if crm else "", crm["stage"].title() if crm else "",
            crm["contacts"] if crm else "", crm["last"] if crm else "",
            CRM_URL.format(crm["ids"][0]) if crm else "", action,
            g("contact_name"), g("contact_title"), g("contact_linkedin_url"), g("public_email"),
            g("public_phone"), src, g("confidence"), who, angle,
            "Not contacted", "", g("notes")])

    order = {"A - Hot": 0, "B - Warm": 1, "C - Nurture": 2, "Partner only": 3}
    rows.sort(key=lambda r: (order[r[0]], -r[1], r[2].lower()))

    ws = wb.create_sheet("Consolidated Outreach", 0)
    headers = ["Priority", "Score", "Company", "Category", "What they do", "Website", "India base",
               "# GCC events sponsored", "Best tier seen", "Latest event date",
               "Sponsorship history (event - tier)", "Sponsorship unconfirmed?",
               "In HubSpot?", "CRM owner", "CRM stage", "CRM contacts", "CRM last contacted", "CRM link",
               "Next action",
               "Contact name (public source)", "Designation", "LinkedIn URL",
               "Published email (company/general)", "Published phone", "Contact source URL(s)",
               "Contact confidence", "Who to contact (roles)", "Pitch angle",
               "Status", "Elets owner", "Research notes / next step"]
    widths = [12, 7, 30, 24, 36, 24, 14, 10, 24, 14, 60, 11,
              10, 22, 10, 10, 14, 30, 40,
              24, 30, 34, 28, 18, 40, 11, 36, 50, 16, 14, 40]
    write_table(ws, "Consolidated", headers, rows, widths, link_cols=(6, 18, 22, 25))

    fills = {"A - Hot": "F8CBAD", "B - Warm": "FFE699", "C - Nurture": "DDEBF7", "Partner only": "E7E6E6"}
    from openpyxl.styles import PatternFill, Font
    crm_fill = {"Yes": "E2EFDA", "No": "FCE4D6"}
    for r in range(2, ws.max_row + 1):
        p = ws.cell(r, 1)
        p.fill = PatternFill("solid", start_color=fills[p.value])
        p.font = Font(name="Arial", size=10, bold=True)
        ws.cell(r, 13).fill = PatternFill("solid", start_color=crm_fill[ws.cell(r, 13).value])
        ws.cell(r, 19).font = Font(name="Arial", size=10, bold=True)
        for c in range(29, 32):
            ws.cell(r, c).fill = PatternFill("solid", start_color="FFFFF2")
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"AC2:AC{ws.max_row}")
    ws.freeze_panes = "D2"
    return rows
