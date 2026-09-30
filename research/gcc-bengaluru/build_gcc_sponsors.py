"""Build the Bengaluru GCC events and sponsors workbook.

Data researched on 2026-09-30 from organiser websites, organiser and sponsor
LinkedIn posts, and press releases. Most event sites show sponsors as logos
only, so each sponsor row records how it was verified.

Run: python3 build_gcc_sponsors.py  ->  bengaluru_gcc_events_sponsors.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from outreach import build_outreach

OUT = Path(__file__).with_name("bengaluru_gcc_events_sponsors.xlsx")

# id, event, organiser, edition, date, venue, status, website, sponsors_page, notes
EVENTS = [
    ("E01", "ET Edge GCC Summit 2026 - Bengaluru", "ET Edge (The Times Group)", "11th", "13 Mar 2026",
     "Bengaluru Marriott Hotel Whitefield", "Held", "https://gcc.et-edge.com/2026-bangalore/",
     "https://gcc.et-edge.com/partner-with-us/",
     "Co-powered by Intellion Offices by Tata Realty. Partner wall is logos only; several tiers unnamed."),
    ("E02", "GCC Summit 2026 Bengaluru (Maxpo) - July edition", "Maxpo Exhibitions Pvt. Ltd.", "Inaugural", "15 Jul 2026",
     "Sheraton Grand Bengaluru at Brigade Gateway", "Held", "https://bengalurugccsummit.com/",
     "https://bengalurugccsummit.com/sponsor",
     "KDEM strategic partner. No tiered sponsor list published."),
    ("E03", "GCC Summit 2026 Bengaluru (bengalurugccsummit.com) - November edition", "Maxpo Exhibitions (likely; site names no organiser)",
     "3rd (per site)", "26 Nov 2026", "Sheraton Grand Bengaluru at Brigade Gateway", "Upcoming",
     "https://bengalurugccsummit.com/", "https://bengalurugccsummit.com/conference/partners",
     "Site promotes '30+ sponsors & partners' but has named none yet. Same day as ET GCC Conclave."),
    ("E04", "Dun & Bradstreet GCC Summit 2026 - Bengaluru Edition", "Dun & Bradstreet India", "2nd", "25 Mar 2026",
     "Shangri-La, Bengaluru", "Held", "https://www.dnb.co.in/events/gcc-blr/",
     "https://www.dnb.co.in/events/gcc-blr/", "Presented by JLL, 63SATS Cybertech and Uber for Business."),
    ("E05", "Zinnov Confluence 2026 (India)", "Zinnov", "19th", "19-20 Aug 2026",
     "Sheraton Grand Bengaluru Whitefield", "Held", "https://confluence.zinnov.com/india/",
     "https://confluence.zinnov.com/india/partners/",
     "6,000+ attendees. Includes Zinnov Awards 2026. List probably incomplete (logo wall)."),
    ("E06", "India IT & GCC PPM Summit 2026", "WHY Summits", "1st (probable)", "28-29 Jul 2026",
     "The Chancery Pavilion, Residency Road", "Held", "https://whysummits.com/india-it-gcc-ppm-summit-2026/",
     "https://whysummits.com/india-it-gcc-ppm-summit-2026/", "Partner list matches the website exactly."),
    ("E07", "HFS India Summit 2026 - GCCs, AI & the Global Operating Model", "HFS Research", "-", "10-11 Feb 2026",
     "The Leela Palace, HAL Old Airport Road", "Held", "https://www.hfsevents.com/india-summit-2026/",
     "https://www.hfsevents.com/website/17380/sponsors/", "15 sponsors; tiers found for 5."),
    ("E08", "People Matters GCC Talent Summit 2026 - Bengaluru", "People Matters", "-", "13 Feb 2026",
     "The Leela Bengaluru", "Held", "https://bengaluru.gcctalentsummit.com/",
     "https://bengaluru.gcctalentsummit.com/", "HR/talent focus. Tiers match site headings."),
    ("E09", "SSON Shared Services & GCC Week India 2026", "SSON (IQPC)", "4th", "21-24 Apr 2026",
     "Taj Yeshwantpur, Bengaluru", "Held", "https://www.ssonetwork.com/events-ssoindia",
     "https://www.ssonetwork.com/events-ssoindia/sponsors", "2027 edition: 20-23 Apr 2027, Bengaluru."),
    ("E10", "EY GCC Conclave 2025 - Bengaluru", "EY India", "-", "Nov 2025 (on/before 20 Nov)",
     "Sheraton Grand Bengaluru Whitefield", "Held",
     "https://www.ey.com/en_in/services/consulting/global-capability-centers/gcc-conclave-2025-bengaluru",
     "-", "EY-hosted, invite-only; no sponsors. 250+ leaders from 120 GCCs. No 2026 Bengaluru date yet."),
    ("E11", "GCC Leadership Conclave - Bengaluru", "The Leadership Federation", "1st", "14 May 2025",
     "Novotel Bengaluru Outer Ring Road", "Held", "https://theleadershipfederation.com/page68571919.html",
     "-", "Sponsors from ANI press release (21 May 2025)."),
    ("E12", "GCC Leadership Conclave - Bengaluru", "The Leadership Federation", "3rd", "3 Sep 2025",
     "Novotel Bengaluru Outer Ring Road", "Held",
     "https://theleadershipfederation.com/gccleadershipconclavebengaluru3rdseptember", "-",
     "Sponsors from ANI/PNN press release (8 Sep 2025)."),
    ("E13", "GCC Leadership Conclave - Bengaluru", "The Leadership Federation", "6th", "7-8 Apr 2026",
     "JW Marriott Bengaluru", "Held", "https://gcc.theleadershipfederation.com/blr", "-",
     "Sponsors from ANI/PNN press release (11 Apr 2026)."),
    ("E14", "GCC Leadership Conclave - Bengaluru", "The Leadership Federation", "9th", "9-10 Sep 2026",
     "The Leela Bhartiya City", "Held", "https://gcc.theleadershipfederation.com/blr", "-",
     "700+ delegates. Sponsors from ANI/PNN press release (15-16 Sep 2026)."),
    ("E15", "ET GCC Annual Conclave & Awards 2025 (ETGCCWorld)", "ET B2B / Times Internet", "2nd", "26 Nov 2025",
     "Sheraton Grand Bengaluru Whitefield", "Held",
     "https://gcc.economictimes.indiatimes.com/gcc-conclave-bengaluru", "-",
     "3rd edition set for 26 Nov 2026, Bengaluru (venue TBA, no partners yet)."),
    ("E16", "ETGCCWorld Talent Conclave 2026", "ET B2B / Times Internet", "-", "6 Aug 2026",
     "Sheraton Grand Bengaluru at Brigade Gateway", "Held",
     "https://gcc.economictimes.indiatimes.com/talent-conclave", "-", "CHRO-led."),
    ("E17", "YourStory GCC Summit 2025", "YourStory Media", "1st", "13 Jun 2025",
     "Bengaluru Marriott Hotel Whitefield", "Held", "https://events.yourstory.com/gcc-summit-2025",
     "https://events.yourstory.com/gcc-summit-2025#partners", "No 2026 edition found."),
    ("E18", "Express Computer GCC Conclave 2026", "Express Computer (Indian Express Group)", "2nd", "4 Feb 2026",
     "JW Marriott Hotel Bengaluru", "Held", "https://gcc.expresscomputer.in/",
     "https://gcc.expresscomputer.in/sponsorship.php", "Partners from Express Computer thank-you post."),
    ("E19", "3AI GCC X...Summit 2026", "3AI", "5th", "20 Feb 2026",
     "Radisson Blu Outer Ring Road", "Held", "https://gccxsummit.com/", "https://gccxsummit.com/",
     "Tiers on site: Knowledge, Gold, Premium, Agentic Data Eng., Lanyard, Exhibit (logos only)."),
    ("E20", "GCC XL Summit 2025", "3AI with ANSR", "-", "24 Sep 2025",
     "The Leela Palace", "Held", "https://gccxlsummit.com/", "-",
     "Invite-only, 125+ GCC heads. 2026 edition moved to Hyderabad."),
    ("E21", "GCC Converge Summit & Awards 2026", "The Mainstream (formerly CIO News)", "8th", "24 Jul 2026",
     "Bengaluru (venue not shown)", "Held", "https://themainstream-gccconverge.co.in/8th-edition/", "-",
     "2025 finale was also Bengaluru (21 Nov 2025; NTT DATA report partner)."),
    ("E22", "nasscom GCC 2030 & Beyond - Bengaluru Chapter", "nasscom with Times Techies", "-", "28/29 Jul 2025",
     "Taj MG Road", "Held", "https://nasscom.in/gcc2030andbeyond/bengaluru-chapter.html", "-",
     "CXO evening. Sponsors thanked but not named."),
    ("E23", "SSF Global 15th Annual GCC Conclave & GCC Excellence Awards 2026", "SSF Global (Shared Services Forum)",
     "15th", "10-11 Dec 2026", "Bengaluru (venue TBA)", "Upcoming", "https://ssfglobal.in/gcc-conclave-2026/", "-",
     "No sponsors announced. Quintes Global is SSF's standing knowledge partner."),
]

# event_id, tier, sponsor, verification
V_SITE, V_ORG, V_SPON, V_PR, V_INF = (
    "Organiser website", "Organiser LinkedIn post", "Sponsor's own post/page",
    "Press release", "Inferred - unconfirmed")
SPONSORS = [
    # E01 ET Edge
    ("E01", "Co-powered by", "Intellion Offices by Tata Realty", V_ORG),
    ("E01", "AI Finance Partner (probable)", "CashFlo", V_INF),
    ("E01", "Solution showcase (tier unknown)", "Haworth", V_INF),
    ("E01", "Solution showcase (tier unknown)", "Scry AI", V_INF),
    ("E01", "Solution showcase (tier unknown)", "Equiniti India", V_INF),
    ("E01", "Knowledge/advisory (probable)", "Deloitte", V_INF),
    ("E01", "Media Partner", "Prop News Time", V_ORG),
    ("E01", "Media Partner (probable)", "ISAIL.in", V_INF),
    # E02 Maxpo July
    ("E02", "Strategic Partner", "Karnataka Digital Economy Mission (KDEM)", V_PR),
    ("E02", "In partnership with", "Dept. of Electronics, IT & BT, Govt of Karnataka", V_PR),
    ("E02", "Collaborator (tier unknown)", "State Bank of India", V_ORG),
    ("E02", "Official Media Partner", "Storify News", V_SPON),
    ("E02", "Official Media Partner", "PharmiVon Media LLP", V_SPON),
    ("E02", "Exhibitor", "Skydo", V_ORG),
    # E04 D&B
    ("E04", "Presenting Partner", "JLL India", V_ORG),
    ("E04", "Presenting Partner", "63SATS Cybertech", V_ORG),
    ("E04", "Presenting Partner", "Uber for Business", V_ORG),
    ("E04", "Business Orchestration Partner", "BMC Software", V_ORG),
    ("E04", "Business Orchestration Partner", "Zertain India", V_ORG),
    ("E04", "Co-Partner", "Power Bridge Systems", V_ORG),
    ("E04", "Branding Partner", "Intellect Design Arena", V_ORG),
    ("E04", "Publication Partner (probable)", "Dhruva Advisors", V_INF),
    ("E04", "Publication Partner (probable)", "World Wide Technology", V_INF),
    # E05 Zinnov
    ("E05", "Host State Partner", "Dept. of Electronics, IT & BT, Govt of Karnataka", V_ORG),
    ("E05", "State Partner", "Invest Punjab", V_ORG),
    ("E05", "State Partner", "Invest UP", V_ORG),
    ("E05", "Platinum", "Everpure", V_ORG),
    ("E05", "Platinum", "Draup", V_ORG),
    ("E05", "Premium Partner", "Snowflake", V_ORG),
    ("E05", "Premium Partner", "Quest Global", V_ORG),
    ("E05", "Global Partner", "NatWest Group", V_ORG),
    ("E05", "Gold", "Cyient", V_ORG),
    ("E05", "Gold", "ProHance", V_ORG),
    ("E05", "Gold", "Commvault", V_ORG),
    ("E05", "Gold", "Everforth", V_ORG),
    ("E05", "Silver", "Dell Technologies", V_ORG),
    ("E05", "Silver", "Sony", V_ORG),
    ("E05", "Silver", "Woodside Energy", V_ORG),
    ("E05", "Silver", "IDFC FIRST Bank", V_ORG),
    ("E05", "Workspace Partner", "Awfis", V_SPON),
    # E06 WHY Summits
    ("E06", "Strategic Partner", "Karnataka Digital Economy Mission (KDEM)", V_SITE),
    ("E06", "Platinum", "monday.com", V_SITE),
    ("E06", "Gold", "Planview", V_SITE),
    ("E06", "Gold", "vSaaS Global", V_SITE),
    # E07 HFS
    ("E07", "Platinum", "EY", V_SPON),
    ("E07", "Platinum", "Ascendion", V_SPON),
    ("E07", "Gold", "Hexaware Technologies", V_SPON),
    ("E07", "Silver", "Tech Mahindra", V_SPON),
    ("E07", "Lead sponsor (tier not named)", "Akkodis", V_SPON),
] + [("E07", "Sponsor (tier not found)", s, V_ORG) for s in (
    "Accenture", "Birlasoft", "EdgeVerve", "Genpact", "GlobalLogic", "HCLTech",
    "IBM", "Infosys", "OrbitShift.AI", "VentureSoft")] + [
    # E08 People Matters
    ("E08", "Presenting Partner", "Flipspaces", V_ORG),
] + [("E08", "Gold Partner", s, V_ORG) for s in (
    "cult.fit", "Udemy Business", "Careernet", "Mandaala", "Klay Preschools & Daycare", "Tortoise")] + [
    ("E08", "Exhibitor", s, V_ORG) for s in (
    "Benepik", "Ceipal", "Shework.in", "Zyla Health", "Silver Brook Learning Center",
    "Syncortex", "Zeko AI", "Mancer Consulting Services")] + [
    # E09 SSON
    ("E09", "Thought Leadership Sponsor", "IBM", V_ORG),
] + [("E09", "Sponsor", s, V_ORG) for s in (
    "Thomson Reuters", "ProHance", "Bluecopa", "Regalia Business Parks (Hiranandani)", "FlexTecs")] + [
    ("E09", "Sponsor", "Sigma AVIT Technology Solutions", V_SPON),
    ("E09", "Media Partner", "SSON Research & Analytics", V_ORG),
    ("E09", "Media Partner", "Indian Chamber of International Business (ICIB)", V_ORG),
    # E11 LF May 2025
    ("E11", "Platinum", "Collabera", V_PR),
    ("E11", "Gold", "Kirtane & Pandit Consulting", V_PR),
    ("E11", "Premium", "Weaver and Tidwell India", V_PR),
    ("E11", "Silver", "Freshworks", V_PR),
    ("E11", "Bronze", "Kovaion Consulting", V_PR),
    ("E11", "Bronze", "ACT Enterprise", V_PR),
    # E12 LF Sep 2025
    ("E12", "Gold", "Cognixia", V_PR),
    ("E12", "Silver", "pCloudy", V_PR),
    ("E12", "Bronze", "BINDZ Consulting", V_PR),
    ("E12", "Gourmet Partner", "Elior India", V_PR),
    # E13 LF Apr 2026
    ("E13", "Gold Partner", "ColorTokens", V_PR),
    ("E13", "Silver Partner", "DevRev", V_PR),
    ("E13", "Ecosystem Partner", "GIFT City", V_PR),
    ("E13", "Academic Partner", "BITSoM", V_PR),
    ("E13", "CXO Dinner Partner", "ADP", V_PR),
    ("E13", "CXO Dinner Partner", "Amicorp", V_PR),
    ("E13", "CXO Dinner Partner", "HSAG Consulting", V_PR),
    ("E13", "Banking Partner", "RBL Bank", V_PR),
    ("E13", "GCC Talent Partner", "Orcapod", V_PR),
    ("E13", "Travel Partner", "Giriraj Mobility", V_PR),
] + [("E13", "Branding Partner", s, V_PR) for s in (
    "E-Solutions", "CoreStack", "Progress", "RST Solutions", "Silver Oak Health",
    "EduTech", "Vidushi Infotech", "Peregrine")] + [
    # E14 LF Sep 2026
    ("E14", "Gold Partner", "Orane Consulting", V_PR),
    ("E14", "Gold Partner", "Axentia.AI", V_PR),
    ("E14", "CXO Round Table Partner", "Brilyant", V_PR),
    ("E14", "Digital Transformation Partner", "Tavant Technologies", V_PR),
    ("E14", "GCC Real Estate Partner", "Regalia Business Parks (Hiranandani)", V_PR),
] + [("E14", "Silver Partner", s, V_PR) for s in (
    "Simple.AI", "CBIZ India", "Hosted.AI", "Accion Labs")] + [
    ("E14", "GCC Strategic Knowledge Partner", "Alliance University", V_PR),
    ("E14", "GCC Staffing & Talent Partner", "Spectrum Talent Management", V_PR),
    ("E14", "Media Partner", "The Sunny Shah Show", V_PR),
] + [("E14", "Branding Partner", s, V_PR) for s in (
    "Avantra", "Famli", "Incture Technologies", "STL Digital", "3R Infotech", "Ideya Labs",
    "Jabra", "Lex Visas", "Infogini Consulting", "Skydo", "Cateworld", "Ingenious",
    "Fieldmark Advisory", "NeoSOFT")] + [
    # E15 ET GCC Conclave 2025
    ("E15", "Presenting Sponsor", "PwC India", V_SPON),
    ("E15", "Gold Partner", "Parkar", V_SPON),
    ("E15", "Exhibitor (tier unknown)", "R Systems", V_SPON),
    ("E15", "Exhibitor (tier unknown)", "TEKsystems", V_SPON),
    # E16 ET Talent Conclave
    ("E16", "Gold", "LinkedIn", V_ORG),
    ("E16", "Gold", "Baker Tilly One India", V_ORG),
    ("E16", "Associate", "Technogen India", V_ORG),
    ("E16", "Co-Associate", "TransformEdge Enterprise Services", V_ORG),
    ("E16", "Supporting", "MAXHUB", V_ORG),
    # E17 YourStory
    ("E17", "Co-Presenting", "Google Cloud", V_SITE),
    ("E17", "Co-Presenting", "Snowflake", V_SITE),
    ("E17", "Workspace Partner", "JLL", V_SITE),
    ("E17", "Communications Partner", "Kaizzen", V_SITE),
    ("E17", "Talent Partner", "Xpheno", V_SITE),
    ("E17", "Community Partner", "Koot", V_SITE),
] + [("E17", "Partner", s, V_SITE) for s in (
    "upGrad", "Publicis Sapient", "WeWork India", "Fourth Partner Energy")] + [
    # E18 Express Computer
    ("E18", "Platinum", "Sheeltron Digital Systems", V_ORG),
    ("E18", "Platinum", "Hewlett Packard Enterprise", V_ORG),
    ("E18", "Agentic AI Partner", "OutSystems", V_ORG),
    ("E18", "Strategic Partner", "K-Tech Innovation Hub (IIIT Dharwad Research Park)", V_ORG),
    ("E18", "Exhibit Partner", "EPAM Systems", V_ORG),
    ("E18", "Associate", "Avekshaa Technologies", V_ORG),
    ("E18", "Associate", "Quess Corp", V_ORG),
    ("E18", "Associate", "Paramatrix Technologies", V_ORG),
    ("E18", "Digital Transformation Partner", "NeoSOFT", V_ORG),
    ("E18", "Workspace Partner", "Table Space", V_ORG),
    # E19 3AI
    ("E19", "Lanyard Partner", "AuthBridge", V_SPON),
    # E20 GCC XL
    ("E20", "Co-host", "ANSR", V_PR),
    ("E20", "Partner (tier unknown)", "C5i", V_SPON),
    # E21 GCC Converge
    ("E21", "Platinum", "Lenovo", V_ORG),
    ("E21", "Platinum", "AMD", V_ORG),
    ("E21", "Gold / AI-Powered TPRM Partner", "UpGuard", V_ORG),
    ("E21", "Gold", "Scry AI", V_ORG),
    ("E21", "Tech Talent Partner", "Quess IT Staffing", V_ORG),
    # E22 nasscom
    ("E22", "Collaboration Partner", "Times Techies (Times of India)", V_SITE),
]

# Normalise company names that appear under slightly different spellings,
# so the repeat-sponsor count is right.
CANON = {"Quess IT Staffing": "Quess Corp", "JLL India": "JLL",
         "Dept. of Electronics, IT & BT, Govt of Karnataka": "Govt of Karnataka (DoITBT)"}

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="D0D0D0")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_table(ws, name, headers, rows, widths, link_cols=()):
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = HEAD_FONT, HEAD_FILL, Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.alignment, cell.border = BODY, WRAP, GRID
            if cell.column in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink, cell.font = cell.value, LINK
    ws.freeze_panes = "B2"
    ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def main():
    ev = {e[0]: e for e in EVENTS}
    wb = Workbook()

    ws = wb.active
    ws.title = "Events"
    n_ev = len(EVENTS)
    rows = [e[:9] + (None,) + e[9:] for e in EVENTS]
    write_table(ws, "Events", ["ID", "Event", "Organiser", "Edition", "Date", "Venue (Bengaluru)", "Status",
                               "Official website", "Sponsors / partners page", "# sponsors found", "Notes"],
                rows, [6, 42, 28, 11, 16, 30, 10, 45, 45, 11, 50], link_cols=(8, 9))
    per_event = {e[0]: sum(1 for sp in SPONSORS if sp[0] == e[0]) for e in EVENTS}
    for i in range(2, n_ev + 2):
        c = ws.cell(i, 10, per_event[ws.cell(i, 1).value])
        c.alignment = Alignment(horizontal="center", vertical="top")

    ws = wb.create_sheet("Sponsors")
    rows = [(eid, ev[eid][1], ev[eid][4], tier, name, CANON.get(name, name), ver)
            for eid, tier, name, ver in SPONSORS]
    write_table(ws, "Sponsors", ["Event ID", "Event", "Date", "Tier (as published)", "Sponsor / partner",
                                 "Company (normalised)", "How verified"],
                rows, [9, 42, 16, 32, 40, 36, 24])
    n_sp = len(rows)

    ws = wb.create_sheet("Repeat Sponsors")
    companies = {}
    for _, tier, name, _ in SPONSORS:
        c = CANON.get(name, name)
        companies.setdefault(c, set())
    counts = {c: len({e for e, _, n, _ in SPONSORS if CANON.get(n, n) == c}) for c in companies}
    ordered = sorted(companies, key=lambda c: (-counts[c], c.lower()))
    hdr = ["Company", "# events sponsored", "Events"]
    rows = []
    for c in ordered:
        evs = sorted({e for e, _, n, _ in SPONSORS if CANON.get(n, n) == c})
        rows.append((c, counts[c], "; ".join(f"{ev[e][1]} ({ev[e][4]})" for e in evs)))
    write_table(ws, "RepeatSponsors", hdr, rows, [40, 12, 110])
    for i in range(2, len(rows) + 2):
        ws.cell(i, 2).alignment = Alignment(horizontal="center", vertical="top")

    outreach = build_outreach(wb, EVENTS, SPONSORS, CANON, write_table)
    wb.active = 0

    ws = wb.create_sheet("Notes")
    notes = [
        "Bengaluru GCC events and sponsors - researched 30 Sep 2026",
        "",
        "Scope: GCC / Global Capability Centre conferences held in Bengaluru, Jan 2025 to Dec 2026 (held and upcoming).",
        "Most event websites show sponsors as logo images only. Names were rebuilt from organiser and sponsor "
        "LinkedIn posts, organiser thank-you posts and ANI/PNN press releases.",
        "The 'How verified' column on the Sponsors sheet says where each name came from. Rows marked "
        "'Inferred - unconfirmed' or with '(probable)' / '(tier unknown)' should be checked before use.",
        "Lists for Zinnov Confluence, ET Edge, HFS and 3AI are probably incomplete; open the sponsors page in a browser "
        "to check the logo wall.",
        "'# sponsors found' and '# events sponsored' are fixed counts from the build script; re-run "
        "build_gcc_sponsors.py after editing the data to refresh them.",
        "",
        "OUTREACH SHEET - how to use:",
        "Priority: A - Hot / B - Warm / C - Nurture, from Score = best tier (Presenting/Platinum 5, Gold/Strategic 3, "
        "Silver/Associate 2, other 1) + 3 per extra GCC event sponsored + 2 if they sponsored in 2026 - 2 if the "
        "sponsorship is unconfirmed. 'Partner only' = government, media or academic partners (in-kind, not paid).",
        "Category, 'Who to contact' and 'Pitch angle' are suggestions, not researched facts.",
        "Contact columns were filled by free web research only (company sites, press releases, LinkedIn posts, "
        "event pages) - no Lusha or Clay credits. Every named contact has a source URL; check the person is still "
        "in the role before writing. Emails/phones are only ones the company publishes (often general or media "
        "lines) - none were guessed. Personal emails and direct phones need Lusha/Clay (credits).",
        "Pale-yellow columns are for your team: update Status (dropdown), Owner and notes. Example: Status "
        "'Contacted', Owner 'Anurag', note 'Sent deck 2 Oct, follow up 9 Oct'.",
        "",
        "Not included (not Bengaluru or not GCC-specific): Nasscom GCC Summit 2025 (Hyderabad) and 2026 (Mumbai); "
        "GCC XL Summit 2026 (Hyderabad); MachineCon GCC (Goa); ETGCCWorld SURGE (Kerala); AIGCC '26 (Chennai); "
        "Bengaluru Tech Summit 2025 (GCC panels only).",
        "Leads without sponsor data: India GCC Conclave 2026 by The Builders Club (Bengaluru, date TBA) - "
        "https://thebuildersclub.me/gcccon/ ; GCC Pulse roundtables - https://gcc-pulse.com/events/ ; "
        "ET GCC Leadership Excellence Masterclass, 19-20 Nov 2026 (paid course).",
    ]
    ws.column_dimensions["A"].width = 130
    for i, t in enumerate(notes, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=FONT, size=12 if i == 1 else 10, bold=(i == 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print(f"{OUT}: {len(outreach)} outreach rows, {n_ev} events, {n_sp} sponsor rows, {len(ordered)} companies")


if __name__ == "__main__":
    main()
